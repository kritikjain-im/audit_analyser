import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

# Define color hexes
HEX_GREEN = "C6EFCE"        # Light green for Match 1
HEX_RED = "FFC7CE"          # Light red for Match 0
HEX_YELLOW = "FFEB9C"       # Light yellow for Accuracy 75-89

font_normal = Font(name="Segoe UI", size=10)
font_bold = Font(name="Segoe UI", size=10, bold=True)
font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Slate Blue
fill_summary = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Light Gold

fill_green = PatternFill(start_color=HEX_GREEN, end_color=HEX_GREEN, fill_type="solid")
fill_red = PatternFill(start_color=HEX_RED, end_color=HEX_RED, fill_type="solid")
fill_yellow = PatternFill(start_color=HEX_YELLOW, end_color=HEX_YELLOW, fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

def auto_fit_columns(ws, min_width=12):
    """Adjust column widths dynamically."""
    for col in ws.columns:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_width)

def create_excel_report(data_calls, field_accuracies, overall_accuracy, output_path):
    wb = openpyxl.Workbook()
    
    # --- Sheet 1: Comparative Analysis ---
    ws = wb.active
    ws.title = "Comparative Analysis"
    
    headers = [
        "Call ID", "Buyer Name", "Buyer Location", "Buyer Contact Details",
        "Lead Tag", "Product Name", "Specifications", "Buyer Required Quantity",
        "Price Quoted by Seller", "Seller Actionables", "Buyer Intent",
        "Buyer Questions", "Reminder", "Matches", "Total Fields", "Accuracy %"
    ]
    ws.append(headers)
    
    # 1. Append Call Rows
    for item in data_calls:
        ws.append([
            item["call_id"],
            item["matches_by_field"]["buyer_name"],
            item["matches_by_field"]["buyer_location"],
            item["matches_by_field"]["buyer_contact"],
            item["matches_by_field"]["lead_tag"],
            item["matches_by_field"]["product_name"],
            item["matches_by_field"]["specifications"],
            item["matches_by_field"]["quantity"],
            item["matches_by_field"]["price"],
            item["matches_by_field"]["actionables"],
            item["matches_by_field"]["buyer_intent"],
            item["matches_by_field"]["buyer_questions"],
            item["matches_by_field"]["reminder"],
            item["matches"],
            item["total_fields"],
            round(item["accuracy"], 2)
        ])
        
    # 2. Append Summary Row
    summary_row = ["Overall Statistics"]
    # Add field accuracies (columns 2 to 13)
    for field_name in [
        "buyer_name", "buyer_location", "buyer_contact", "lead_tag",
        "product_name", "specifications", "quantity", "price",
        "actionables", "buyer_intent", "buyer_questions", "reminder"
    ]:
        acc = field_accuracies.get(field_name, 0.0)
        summary_row.append(f"{acc:.2f}%")
        
    # Add trailing columns
    summary_row.append("") # Matches column blank
    summary_row.append("") # Total Fields column blank
    summary_row.append(f"Overall Accuracy: {overall_accuracy:.2f}%")
    
    ws.append(summary_row)
    
    # 3. Base Formatting
    ws.freeze_panes = "A2"
    
    # Headers styling
    for cell in ws[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    # Data row styling
    for r_idx in range(2, ws.max_row):
        for cell in ws[r_idx]:
            cell.font = font_normal
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    # Bottom Summary Row styling
    summary_row_idx = ws.max_row
    for cell in ws[summary_row_idx]:
        cell.font = font_bold
        cell.fill = fill_summary
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    # 4. Conditional Formatting for Match columns (Columns B to M, i.e., 2 to 13)
    match_rule_1 = CellIsRule(operator='equal', formula=['1'], fill=fill_green)
    match_rule_0 = CellIsRule(operator='equal', formula=['0'], fill=fill_red)
    for col_idx in range(2, 14):
        col_letter = get_column_letter(col_idx)
        # Apply only to data rows
        ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{summary_row_idx-1}", match_rule_1)
        ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{summary_row_idx-1}", match_rule_0)
        
    # Conditional formatting for Accuracy column (Column P / 16)
    acc_rule_green = CellIsRule(operator='greaterThanOrEqual', formula=['90'], fill=fill_green)
    acc_rule_yellow = CellIsRule(operator='between', formula=['75', '89.9'], fill=fill_yellow)
    acc_rule_red = CellIsRule(operator='lessThan', formula=['75'], fill=fill_red)
    ws.conditional_formatting.add(f"P2:P{summary_row_idx-1}", acc_rule_green)
    ws.conditional_formatting.add(f"P2:P{summary_row_idx-1}", acc_rule_yellow)
    ws.conditional_formatting.add(f"P2:P{summary_row_idx-1}", acc_rule_red)
    
    auto_fit_columns(ws)

    # --- Sheet 2: Detailed Value Comparison ---
    ws2 = wb.create_sheet(title="Detailed Value Comparison")
    
    field_details = [
        ("buyer_name", "Buyer Name"),
        ("buyer_location", "Buyer Location"),
        ("buyer_contact", "Buyer Contact Details"),
        ("lead_tag", "Lead Tag"),
        ("product_name", "Product Name"),
        ("specifications", "Specifications"),
        ("quantity", "Buyer Required Quantity"),
        ("price", "Price Quoted by Seller"),
        ("actionables", "Seller Actionables"),
        ("buyer_intent", "Buyer Intent"),
        ("buyer_questions", "Buyer Questions"),
        ("reminder", "Reminder")
    ]
    
    # 25 Columns: Call ID + (12 fields * 2)
    headers2 = ["Call ID"]
    for _, field_label in field_details:
        headers2.append(f"{field_label} (M)")
        headers2.append(f"{field_label} (AI)")
    ws2.append(headers2)
    
    for item in data_calls:
        row_data = [item["call_id"]]
        for field_key, _ in field_details:
            m_val = item["manual_values"].get(field_key)
            a_val = item["ai_values"].get(field_key)
            
            # Format/clean output representation
            row_data.append(m_val if m_val is not None else "")
            row_data.append(a_val if a_val is not None else "")
        ws2.append(row_data)
        
    # Styles for Sheet 2
    ws2.freeze_panes = "A2"
    
    for cell in ws2[1]:
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    for r_idx in range(2, ws2.max_row + 1):
        # First column (Call ID) is bold
        ws2.cell(row=r_idx, column=1).font = font_bold
        for cell in ws2[r_idx]:
            if cell.column > 1:
                cell.font = font_normal
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    auto_fit_columns(ws2)
    
    # --- Sheet 3: Summary Statistics ---
    ws3 = wb.create_sheet(title="Summary Statistics")
    
    # Updated headers to include "Inferred"
    headers3 = [
        "Field Name", "Correct", "Correct (Not Discussed)", "Inferred", "Incorrect",
        "Missing", "Hallucination", "Matched Calls", "Total calls", "Accuracy"
    ]
    ws3.append(headers3)
    
    total_calls = len(data_calls)
    field_rows = []
    
    for field_key, field_label in field_details:
        counts = {
            "Correct": 0,
            "Correct (Not Discussed)": 0,
            "Inferred": 0,
            "Incorrect": 0,
            "Missing": 0,
            "Hallucination": 0
        }
        
        matched_calls = 0
        for item in data_calls:
            val = item["manual_values"].get(field_key)
            if val in counts:
                counts[val] += 1
            matched_calls += item["matches_by_field"].get(field_key, 0)
            
        acc = (matched_calls / total_calls) * 100 if total_calls > 0 else 0.0
        
        row_data = [
            field_label,
            counts["Correct"],
            counts["Correct (Not Discussed)"],
            counts["Inferred"],
            counts["Incorrect"],
            counts["Missing"],
            counts["Hallucination"],
            matched_calls,
            total_calls,
            acc
        ]
        ws3.append(row_data)
        field_rows.append(acc)
        
    # Overall summary row (10 columns)
    overall_acc = sum(field_rows) / len(field_rows) if field_rows else 0.0
    ws3.append([
        "Overall Accuracy:", "", "", "", "", "", "", "", "", overall_acc
    ])
    
    # Header 3 styling
    fill_header3 = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid") # Soft Ice Blue
    font_bold_10 = Font(name="Segoe UI", size=10, bold=True)
    
    for cell in ws3[1]:
        cell.font = font_bold_10
        cell.fill = fill_header3
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        
    # Data rows styling
    for r_idx in range(2, ws3.max_row):
        ws3.cell(row=r_idx, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws3.cell(row=r_idx, column=1).font = font_normal
        ws3.cell(row=r_idx, column=1).border = thin_border
        
        for c_idx in range(2, 10):
            cell = ws3.cell(row=r_idx, column=c_idx)
            cell.font = font_normal
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
        ws3.cell(row=r_idx, column=10).number_format = '0"%"'
        
    # Overall summary row styling
    summary_r_idx = ws3.max_row
    ws3.cell(row=summary_r_idx, column=1).alignment = Alignment(horizontal="left", vertical="center")
    
    for cell in ws3[summary_r_idx]:
        cell.font = font_bold_10
        cell.fill = fill_summary
        cell.border = thin_border
        if cell.column > 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    ws3.cell(row=summary_r_idx, column=10).number_format = '0.00"%"'
    
    ws3.freeze_panes = "A2"
    auto_fit_columns(ws3)
    
    wb.save(output_path)
